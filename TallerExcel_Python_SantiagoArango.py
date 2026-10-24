# Taller Python + Excel - Santiago Arango
import openpyxl                  # para leer archivos Excel
from tabulate import tabulate    # para mostrar datos en forma de tabla

# Abrir el archivo Excel (si no existe, avisa y termina)
try:
    excel_dataframe = openpyxl.load_workbook("personas.xlsx")
except FileNotFoundError:
    print("No se encontro el archivo personas.xlsx")
    exit()

dataframe = excel_dataframe.active   # hoja activa (donde estan los datos)

data = []   # aqui se guardan todas las filas del Excel

# Recorrer filas desde la 2 (la 1 son los titulos)
for row in range(2, dataframe.max_row + 1):
    _row = [row - 1]   # fila temporal, empieza con el numero consecutivo

    # Recorrer columnas y guardar cada valor de la fila
    for col in dataframe.iter_cols(1, dataframe.max_column):
        _row.append(col[row - 1].value)

    data.append(_row)   # guardar la fila completa

# Titulos de la tabla
headers = ["#", "Id", "Name", "Company", "Email", "MAC Address"]
headers_align = ("center",) * 6


# Muestra una lista en tabla (o un aviso si esta vacia)
def mostrar_tabla(lista, titulos, mensaje_vacio):
    if len(lista) == 0:
        print(mensaje_vacio)
    else:
        print(tabulate(lista, headers=titulos, tablefmt="fancy_grid", colalign=("center",) * len(titulos)))


# Opcion 1: ver todos los registros
def ver_todos():
    mostrar_tabla(data, headers, "No hay registros")


# Reto 1: solo Name y Email
def name_email():
    lista = []
    for persona in data:
        lista.append([persona[0], persona[2], persona[4]])
    mostrar_tabla(lista, ["#", "Name", "Email"], "No hay registros")


# Reto 2: contar registros
def contar():
    print("El archivo tiene", len(data), "registros")


# Reto 4: buscar una persona por ID
def buscar_id():
    id_buscado = input("Ingrese el ID: ").strip()
    if id_buscado == "":
        print("Debe ingresar un ID")
        return
    lista = []
    for persona in data:
        if str(persona[1]).lower() == id_buscado.lower():
            lista.append(persona)
    mostrar_tabla(lista, headers, "No se encontro ninguna persona con ese ID")


# Reto 3: filtrar por palabra en la empresa (Enter = "Group")
def filtrar_empresa():
    palabra = input("Palabra a buscar en la empresa (Enter = Group): ").strip()
    if palabra == "":
        palabra = "Group"
    lista = []
    for persona in data:
        if palabra.lower() in str(persona[3]).lower():
            lista.append(persona)
    mostrar_tabla(lista, headers, "No hay empresas con la palabra " + palabra)


# Reto 5: menu interactivo (se repite hasta elegir 6)
def menu():
    opcion = ""
    while opcion != "6":
        print()
        print("========= MENU =========")
        print("1. Ver todos")
        print("2. Buscar por ID")
        print("3. Filtrar empresa")
        print("4. Contar registros")
        print("5. Mostrar Name y Email")
        print("6. Salir")
        opcion = input("Elija una opcion: ").strip()
        print()

        if opcion == "1":
            ver_todos()
        elif opcion == "2":
            buscar_id()
        elif opcion == "3":
            filtrar_empresa()
        elif opcion == "4":
            contar()
        elif opcion == "5":
            name_email()
        elif opcion == "6":
            print("Hasta luego")
        else:
            print("Opcion no valida, escriba un numero del 1 al 6")


menu()   # arranca el programa
